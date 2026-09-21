"""Extract both measured grating ranges; read-only source workbook audit.

Run with bundled Python (openpyxl). No change to the legacy spectral dataset.
"""
from pathlib import Path
import hashlib
import io
import json
import zipfile
import openpyxl
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
MEMBER = '41598_2020_74742_MOESM2_ESM.xlsx'
EXPECTED_SHA = '740f8517fc0f6ce18d724d71568db08d2c58567551057a0008068f8f88de8044'


def extract(archive):
    with zipfile.ZipFile(archive) as z:
        raw = z.read(MEMBER)
    if hashlib.sha256(raw).hexdigest() != EXPECTED_SHA:
        raise ValueError('Source workbook hash changed; review before extraction')
    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheet = workbook['Sheet 1']
    bands = []
    for label, start, stop, columns, channels, limits in [
        ('short', 3, 50, (4, 15), ['Rh1', 'Rh3', 'Rh4', 'Rh5', 'Rh6'], (315, 550)),
        ('long', 3, 53, (15, 20), ['Rh1', 'Rh6'], (450, 700)),
    ]:
        if [sheet.cell(1, c).value for c in range(columns[0]+1, columns[1], 2)] != channels:
            raise ValueError('Unexpected source column labels')
        values = np.array([[sheet.cell(r, c).value for c in range(*columns)]
                           for r in range(start, stop+1)], dtype=float)
        if not np.array_equal(values[:, 0], np.arange(limits[0], limits[1]+1, 5)):
            raise ValueError('Unexpected wavelength grid')
        if not np.isfinite(values).all() or np.any(values[:, 1:] < 0):
            raise ValueError('Invalid measured values')
        bands.append(dict(name=label, channels=channels, wavelength_nm=values[:, 0].tolist(),
                          mean=values[:, 1::2].tolist(), sd=values[:, 2::2].tolist(), n=6,
                          cells='D3:N50' if label == 'short' else 'O3:S53'))
    workbook.close()
    return dict(source='https://www.nature.com/articles/s41598-020-74742-1',
        download='https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7588446/supplementaryFiles',
        source_member=MEMBER, source_sha256=EXPECTED_SHA, sheet='Sheet 1',
        measurement='Red-eye single-opsin-rescue ERG mean and sample SD; n=6 per band',
        bands=bands, rh1_long_scale_from_methods=0.5777,
        limitations=['ERG shapes are not absolute single-cell photon catch or voltage calibration.',
                     'No Rh3/Rh4/Rh5 measurements above 550 nm in these source columns.',
                     'No individual animal traces or temporal responses in these columns.'])


if __name__ == '__main__':
    result = extract(ROOT/'data/spectral/supplementary.zip')
    out = ROOT/'research/calibration/sharkey2020_bands.json'
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2)+'\n')
    print(out)
